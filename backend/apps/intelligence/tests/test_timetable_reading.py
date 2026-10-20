from datetime import date, time
from io import BytesIO

import inngest
import openpyxl
import pytest
from django.conf import settings
from django.core.files.base import ContentFile
from rest_framework.exceptions import PermissionDenied, ValidationError

from apps.academics.models import Intake, Programme, School, Term
from apps.accounts import semesters
from apps.accounts.services import create_user
from apps.intelligence import functions
from apps.intelligence.models import MessageFile, TimetableImport
from apps.intelligence.providers import AIError, AIRateLimited, fake_provider
from apps.intelligence.timetable_reading import (
    best_tab,
    confirm_import,
    discard_import,
    read_file,
    spreadsheet_tabs,
    start_import,
)
from apps.timetable.models import CourseSection, SectionEnrollment

pytestmark = pytest.mark.django_db

BCS_CLASSES = {
    "classes": [
        {
            "course_code": "CDE2313",
            "course_name": "Data Structures",
            "section": "A",
            "meetings": [{"day": "monday", "start": "10:00", "end": "12:00", "room": "B2"}],
        }
    ]
}
BBA_CLASSES = {
    "classes": [
        {
            "course_code": "MGT1001",
            "course_name": "Management",
            "section": "1",
            "meetings": [{"day": "tuesday", "start": "08:00", "end": "10:00"}],
        }
    ]
}


@pytest.fixture(autouse=True)
def setting(settings, tmp_path, monkeypatch):
    settings.MEDIA_ROOT = tmp_path
    monkeypatch.setattr(semesters, "today", lambda: date(2025, 11, 1))
    Term.objects.create(
        name="2025/2026 Semester 1",
        academic_session="2025/2026",
        number=1,
        starts_on=date(2025, 10, 27),
        ends_on=date(2026, 2, 27),
    )


@pytest.fixture
def school():
    return School.objects.create(code="SCI", name="School of Computing and Informatics")


def make_student(email, school):
    student = create_user(email, "Str0ng-pass-99", full_name=email)
    profile = student.student_profile
    profile.school = school
    profile.programme = Programme.objects.get_or_create(
        code="BCS",
        defaults={
            "school": school,
            "name": "Bachelor in Computer Science (Honours)",
            "level": "bachelor",
        },
    )[0]
    profile.intake = Intake.objects.get_or_create(
        code="2025-OCT", defaults={"name": "October 2025", "start_date": date(2025, 10, 1)}
    )[0]
    profile.save()
    return student


def xlsx_file() -> MessageFile:
    workbook = openpyxl.Workbook()
    first = workbook.active
    first.title = "BBA Mar 2025"
    first.append(["Day", "Time", "Course", "Room"])
    first.append(["Tuesday", time(8), "MGT1001 Management", "A1"])
    second = workbook.create_sheet("BCS Oct 2025")
    second.append(["Day", "Time", "Course", "Room"])
    second.append(["Monday", time(10), "CDE2313 Data Structures (A)", "B2"])
    workbook.create_sheet("Notes")  # empty: skipped
    buffer = BytesIO()
    workbook.save(buffer)
    return stored(buffer.getvalue(), "spreadsheet", "application/vnd.ms-excel", "xlsx")


def stored(data, kind, mime, extension) -> MessageFile:
    item = MessageFile(kind=kind, mime_type=mime, size_bytes=len(data))
    item.file.save(f"file.{extension}", ContentFile(data), save=False)
    item.save()
    return item


# --- Reading a file once ------------------------------------------------------------------


def test_each_tab_becomes_text():
    tabs = spreadsheet_tabs(xlsx_file().file.open("rb").read())
    assert list(tabs) == ["BBA Mar 2025", "BCS Oct 2025"]  # the empty tab is left out
    assert "Monday | 10:00 | CDE2313 Data Structures (A) | B2" in tabs["BCS Oct 2025"]


def test_a_spreadsheet_is_read_tab_by_tab_only_once():
    file = xlsx_file()
    fake_provider().answers.extend([BBA_CLASSES, BCS_CLASSES])

    tabs = read_file(file)
    assert tabs["BCS Oct 2025"][0]["course_code"] == "CDE2313"
    assert tabs["BBA Mar 2025"][0]["meetings"][0]["day"] == "tuesday"
    calls = fake_provider().calls
    assert [c["model"] for c in calls] == [settings.AI_SMART_MODEL] * 2
    assert 'Timetable tab "BCS Oct 2025"' in calls[1]["prompt"]

    file.refresh_from_db()
    assert read_file(file) == tabs
    assert len(fake_provider().calls) == 2  # the second student costs no AI request


def test_an_image_or_pdf_is_sent_as_an_attachment():
    pdf = stored(b"%PDF-1.7 timetable", "pdf", "application/pdf", "pdf")
    fake_provider().answers.append(BCS_CLASSES)

    assert list(read_file(pdf)) == [""]
    [call] = fake_provider().calls
    assert call["attachments"][0].mime_type == "application/pdf"


# --- Choosing the student's tab -----------------------------------------------------------


def test_the_tab_matching_programme_and_intake_is_chosen(school):
    profile = make_student("ali@student.aiu.edu.my", school).student_profile
    assert best_tab(["BBA Mar 2025", "BCS Oct 2025", "BCS Mar 2025"], profile) == "BCS Oct 2025"
    assert best_tab(["Sheet1", "Sheet2"], profile) == "Sheet1"  # no clue: the first tab


# --- One student's import -----------------------------------------------------------------


def ready_import(student, file, answers=(BBA_CLASSES, BCS_CLASSES)):
    fake_provider().answers.extend(answers)
    timetable_import = start_import(student, file)
    functions.read_timetable_import(timetable_import.id)
    timetable_import.refresh_from_db()
    return timetable_import


def test_read_check_and_confirm_my_timetable(school):
    ali = make_student("ali@student.aiu.edu.my", school)
    timetable_import = ready_import(ali, xlsx_file())

    assert timetable_import.status == "ready"
    assert timetable_import.tab == "BCS Oct 2025"
    assert [row["course_code"] for row in timetable_import.rows] == ["CDE2313"]
    assert CourseSection.objects.count() == 0  # nothing saved before confirming

    confirm_import(ali, timetable_import)
    timetable_import.refresh_from_db()
    assert timetable_import.status == "confirmed"
    assert timetable_import.report["created_sections"] == ["CDE2313 Section A"]
    section = CourseSection.objects.get()
    assert SectionEnrollment.objects.filter(user=ali, section=section).exists()


def test_a_classmate_reuses_the_reading_and_the_sections(school):
    file = xlsx_file()
    confirm_import(ali := make_student("ali@student.aiu.edu.my", school), ready_import(ali, file))
    sara = make_student("sara@student.aiu.edu.my", school)

    sara_import = ready_import(sara, file, answers=())  # no AI answers needed
    confirm_import(sara, sara_import)

    assert CourseSection.objects.count() == 1
    assert SectionEnrollment.objects.filter(section__code="A").count() == 2


def test_the_student_can_fix_the_rows_before_confirming(school):
    ali = make_student("ali@student.aiu.edu.my", school)
    timetable_import = ready_import(ali, xlsx_file())
    fixed = [dict(timetable_import.rows[0], section="B")]

    confirm_import(ali, timetable_import, rows=fixed)
    assert CourseSection.objects.get().code == "B"


def test_only_the_owner_can_confirm_once_and_discard(school):
    ali = make_student("ali@student.aiu.edu.my", school)
    sara = make_student("sara@student.aiu.edu.my", school)
    timetable_import = ready_import(ali, xlsx_file())

    with pytest.raises(PermissionDenied):
        confirm_import(sara, timetable_import)
    discard_import(ali, timetable_import)
    assert timetable_import.status == "discarded"
    with pytest.raises(ValidationError):
        confirm_import(ali, timetable_import)


def test_problems_are_explained_or_retried(school):
    ali = make_student("ali@student.aiu.edu.my", school)

    nothing = ready_import(ali, xlsx_file(), answers=({"classes": []}, {"classes": []}))
    assert (nothing.status, nothing.error) == ("failed", "No classes were found in this file.")

    broken = ready_import(ali, xlsx_file(), answers=(AIError("bad key"),))
    assert broken.status == "failed"
    assert "try again later" in broken.error

    fake_provider().answers.append(AIRateLimited("429"))
    waiting = start_import(ali, xlsx_file())
    with pytest.raises(inngest.RetryAfterError):
        functions.read_timetable_import(waiting.id)
    assert TimetableImport.objects.get(pk=waiting.pk).status == "reading"
